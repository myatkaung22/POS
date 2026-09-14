import json
import re
from pathlib import Path

import openpyxl

SRC = Path(r"c:\Users\myatk\Desktop\POS\Menu\Kitchen Menu list.xlsx")
OUT = Path(r"c:\Users\myatk\Desktop\POS\prisma\kitchen-menu.json")

EMOJI = {
    "Main Dish": "🍽️",
    "Appetizer": "🥟",
    "Salad": "🥗",
    "Grilled (Chicken)": "🍗",
    "Grilled (Pork)": "🥓",
    "Grilled (Beef)": "🥩",
    "Grilled (Seafood)": "🦐",
    "Grilled (Vegetables)": "🥦",
    "Soup": "🍲",
    "Mala": "🌶️",
    "Classic Cocktails": "🍹",
    "Beer Bottle": "🍺",
    "Draft Beer": "🍻",
    "Soft Drinks": "🥤",
    "Brandy": "🥃",
    "Gin": "🍸",
    "Rum": "🥃",
    "Tequila": "🌵",
    "Vodka": "🍸",
    "Whisky": "🥃",
    "Liqueur": "🍷",
}

DRINK_CATS = {
    "Classic Cocktails",
    "Beer Bottle",
    "Draft Beer",
    "Soft Drinks",
    "Brandy",
    "Gin",
    "Rum",
    "Tequila",
    "Vodka",
    "Whisky",
    "Liqueur",
}

FOOD_ORDER = [
    "Appetizer",
    "Salad",
    "Main Dish",
    "Soup",
    "Mala",
    "Grilled (Chicken)",
    "Grilled (Pork)",
    "Grilled (Beef)",
    "Grilled (Seafood)",
    "Grilled (Vegetables)",
]
DRINK_ORDER = [
    "Classic Cocktails",
    "Beer Bottle",
    "Draft Beer",
    "Soft Drinks",
    "Brandy",
    "Gin",
    "Rum",
    "Tequila",
    "Vodka",
    "Whisky",
    "Liqueur",
]


def tidy_part(s):
    return re.sub(r",\s*", ", ", s)


def clean(text):
    if text is None:
        return ""
    s = str(text).replace("\xa0", " ").strip()
    s = s.replace("J�germeister", "Jagermeister").replace("Jägermeister", "Jagermeister")
    s = re.sub(r"\s+", " ", s)
    return s


def parse_prices(val):
    if val is None or val == "":
        return []
    if isinstance(val, (int, float)):
        return [round(float(val), 2)]
    parts = [p.strip() for p in str(val).replace(",", "").split("/") if p.strip()]
    out = []
    for p in parts:
        try:
            out.append(round(float(p), 2))
        except ValueError:
            return []
    return out


def split_named(name, prices):
    name = clean(name)
    m = re.search(r"\(([^)]+/[^)]+)\)", name)
    if m and len(prices) >= 2:
        left, right = [tidy_part(x.strip()) for x in m.group(1).split("/", 1)]
        base = name[: m.start()].strip()
        return [
            (f"{base} ({left})", "", prices[0]),
            (f"{base} ({right})", "", prices[1]),
        ]
    if "/" in name and len(prices) >= 2:
        left, right = [x.strip() for x in name.split("/", 1)]
        words = right.split()
        if len(words) > 1:
            suffix = " " + " ".join(words[1:])
            return [(left + suffix, "", prices[0]), (words[0] + suffix, "", prices[1])]
        return [(left, "", prices[0]), (right, "", prices[1])]
    return None


def variants(name, price_val, remark=""):
    name = clean(name)
    remark = clean(remark)
    prices = parse_prices(price_val)
    if not name or not prices:
        return []
    named = split_named(name, prices)
    if named:
        return named
    remark_parts = [p.strip() for p in remark.split("/") if p.strip()]
    if len(prices) > 1 and len(remark_parts) == len(prices):
        return [(f"{name} ({part})", part, price) for part, price in zip(remark_parts, prices)]
    desc = remark
    return [(name, desc, prices[0])]


def add_item(buckets, category, name, description, price, extra_desc=""):
    category = clean(category)
    if not category:
        return
    desc = " · ".join([p for p in [clean(description), clean(extra_desc)] if p])
    buckets.setdefault(category, []).append(
        {
            "name": clean(name),
            "description": desc,
            "price": price,
        }
    )


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    buckets = {}

    ws = wb["Sheet1"]
    current = None
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=5, values_only=True):
        a, b, _c, _d, e = row
        name = clean(b)
        if not name or name == "Price":
            continue
        if a is None and (e is None or clean(e) == "Price"):
            current = name
            continue
        if current:
            for n, desc, price in variants(name, e):
                add_item(buckets, current, n, desc, price)

    ws = wb["ShotsBottles"]
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=6, values_only=True):
        _no, cat, name, glass, bottle, remark = row
        cat = clean(cat)
        name = clean(name)
        if not cat or not name:
            continue
        glass_prices = parse_prices(glass)
        bottle_prices = parse_prices(bottle)
        remark = clean(remark)
        if glass_prices:
            add_item(buckets, cat, name, "Glass", glass_prices[0])
        if len(bottle_prices) == 1:
            add_item(buckets, cat, f"{name} (Bottle)", remark or "Bottle", bottle_prices[0])
        elif len(bottle_prices) > 1:
            parts = [p.strip() for p in remark.split("/") if p.strip()]
            if len(parts) == len(bottle_prices):
                for part, price in zip(parts, bottle_prices):
                    add_item(buckets, cat, f"{name} (Bottle {part})", part, price)
            else:
                add_item(buckets, cat, f"{name} (Bottle)", remark or "Bottle", bottle_prices[0])

    ws = wb["Cocktails"]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=6, values_only=True):
        cat, _no, name, _desc, ingredients, price = row
        for n, desc, p in variants(name, price, ingredients):
            add_item(buckets, cat or "Classic Cocktails", n, desc or clean(ingredients), p)

    for sheet in ("Beer", "Soft drinks"):
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=5, values_only=True):
            _no, cat, name, price, remark = row
            name = clean(name)
            remark = clean(remark)
            if remark and remark.lower() not in name.lower():
                name = f"{name} ({remark})"
            for n, desc, p in variants(name, price, remark):
                add_item(buckets, cat, n, desc, p)

    prefixes = {}
    def sku_for(cat, index):
        if cat not in prefixes:
            words = re.findall(r"[A-Za-z]+", cat)
            prefixes[cat] = ("".join(w[0] for w in words)[:4] or "MN").upper()
        return f"{prefixes[cat]}-{index:02d}"

    categories = []
    seen = set()
    for cat in FOOD_ORDER + DRINK_ORDER:
        items = buckets.get(cat) or []
        if not items:
            continue
        seen.add(cat)
        unique = []
        used = set()
        for item in items:
            key = (item["name"].lower(), item["price"])
            if key in used:
                continue
            used.add(key)
            unique.append(item)
        categories.append(
            {
                "name": cat,
                "emoji": EMOJI.get(cat, "🍽️"),
                "kitchenPrint": cat not in DRINK_CATS,
                "items": [
                    {
                        **item,
                        "sku": sku_for(cat, i + 1),
                        "emoji": EMOJI.get(cat, "🍽️"),
                        "kitchenPrint": cat not in DRINK_CATS,
                    }
                    for i, item in enumerate(unique)
                ],
            }
        )

    for cat, items in buckets.items():
        if cat in seen:
            continue
        categories.append(
            {
                "name": cat,
                "emoji": EMOJI.get(cat, "🍽️"),
                "kitchenPrint": cat not in DRINK_CATS,
                "items": [
                    {
                        **item,
                        "sku": sku_for(cat, i + 1),
                        "emoji": EMOJI.get(cat, "🍽️"),
                        "kitchenPrint": cat not in DRINK_CATS,
                    }
                    for i, item in enumerate(items)
                ],
            }
        )

    payload = {
        "source": "Menu/Kitchen Menu list.xlsx",
        "categories": categories,
        "itemCount": sum(len(c["items"]) for c in categories),
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print("wrote", OUT, "categories", len(categories), "items", payload["itemCount"])
    for c in categories:
        print(f"  {c['name']}: {len(c['items'])}")


if __name__ == "__main__":
    main()
